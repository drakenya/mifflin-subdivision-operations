from pathlib import Path

import openpyxl

from waybill_generator.importers.roster_xlsx import (
    RawRosterRow, read_freight_cars, is_sold, load_car_type_map, resolve_aar_code,
    extract_tonnage, extract_length_ft, resolve_capacity_tons, CAPACITY_DEFAULTS,
    build_cars, RosterImportReport,
)


def _make_workbook(tmp_path: Path, rows: list[dict[str, object]]) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Freight Cars"
    headers = {
        "A": "Inventory date", "B": "814", "C": "Manufacturer", "D": "Manufacturer #",
        "E": "Location", "F": "Type", "K": "Car #", "S": "Notes",
    }
    for col, val in headers.items():
        ws[f"{col}1"] = val
    for i, row in enumerate(rows, start=2):
        for col, val in row.items():
            ws[f"{col}{i}"] = val
    path = tmp_path / "roster.xlsx"
    wb.save(path)
    return path


class TestReadFreightCars:
    def test_reads_basic_row(self, tmp_path):
        path = _make_workbook(tmp_path, [
            {"A": "45711", "B": "PRR", "E": "Bin 1", "F": "40' Box Car",
             "K": "12345", "S": "a note"},
        ])
        rows = read_freight_cars(path)
        assert len(rows) == 1
        assert rows[0] == RawRosterRow(
            inventory_status="45711", road="PRR", location="Bin 1",
            type_text="40' Box Car", car_number="12345", notes="a note",
        )

    def test_row_with_blank_leading_cells_still_reads_correctly(self, tmp_path):
        # Regression case: Excel omits blank cells from a row entirely, so a
        # positional (non-reference-based) reader would misalign columns here.
        path = _make_workbook(tmp_path, [
            {"B": "ATSF", "F": "40' Box Car", "K": "140185"},
        ])
        rows = read_freight_cars(path)
        assert len(rows) == 1
        assert rows[0].road == "ATSF"
        assert rows[0].type_text == "40' Box Car"
        assert rows[0].car_number == "140185"

    def test_fully_blank_trailing_row_skipped(self, tmp_path):
        path = _make_workbook(tmp_path, [
            {"A": "45711", "B": "PRR", "F": "40' Box Car", "K": "12345"},
            {},
        ])
        rows = read_freight_cars(path)
        assert len(rows) == 1

    def test_numeric_cell_values_stringified(self, tmp_path):
        path = _make_workbook(tmp_path, [
            {"B": "PRR", "F": "40' Box Car", "K": 12345},
        ])
        rows = read_freight_cars(path)
        assert rows[0].car_number == "12345"


def _row(inventory_status="45711", road="PRR", location="Bin 1",
         type_text="40' Box Car", car_number="12345", notes="") -> RawRosterRow:
    return RawRosterRow(
        inventory_status=inventory_status, road=road, location=location,
        type_text=type_text, car_number=car_number, notes=notes,
    )


class TestIsSold:
    def test_sold_in_inventory_status(self):
        assert is_sold(_row(inventory_status="SOLD"))

    def test_sell_in_location(self):
        assert is_sold(_row(location="Sell - Box 02 - Built"))

    def test_listed_swap_in_location(self):
        assert is_sold(_row(location="LISTED - HO SWAP"))

    def test_case_insensitive(self):
        assert is_sold(_row(inventory_status="sold"))

    def test_normal_row_not_sold(self):
        assert not is_sold(_row(inventory_status="45711", location="Bin 08 - Foreign Roads"))

    def test_blank_status_not_sold(self):
        assert not is_sold(_row(inventory_status="", location=""))


class TestLoadCarTypeMap:
    def test_loads_entries_lowercased(self, tmp_path):
        p = tmp_path / "car_type_map.yaml"
        p.write_text(
            "- source_text: \"GS 40' Gondola\"\n  aar_code: GS\n"
            "- source_text: \"Bev-Bel Special Run Box Car\"\n  aar_code: XM\n"
        )
        cmap = load_car_type_map(p)
        assert cmap["gs 40' gondola"] == "GS"
        assert cmap["bev-bel special run box car"] == "XM"

    def test_missing_file_returns_empty(self, tmp_path):
        assert load_car_type_map(tmp_path / "nonexistent.yaml") == {}


class TestResolveAarCode:
    def test_explicit_map_wins(self):
        cmap = {"custom decorated hopper": "LO"}
        code, tier = resolve_aar_code("Custom Decorated Hopper", cmap)
        assert code == "LO"
        assert tier == "map"

    def test_keyword_box_car(self):
        code, tier = resolve_aar_code("40' Steel SD Box Car", {})
        assert code == "XM"
        assert tier == "keyword"

    def test_keyword_open_hopper(self):
        code, tier = resolve_aar_code("70 Ton 12 Panel Triple Hopper", {})
        assert code == "HM"
        assert tier == "keyword"

    def test_keyword_covered_hopper(self):
        code, tier = resolve_aar_code("PS-2 Covered Hopper", {})
        assert code == "LO"
        assert tier == "keyword"

    def test_keyword_grain_hopper(self):
        code, tier = resolve_aar_code("Grain Hopper Car", {})
        assert code == "LB"
        assert tier == "keyword"

    def test_keyword_cement_hopper(self):
        code, tier = resolve_aar_code("Cement Covered Hopper", {})
        assert code == "LC"
        assert tier == "keyword"

    def test_keyword_steel_gondola(self):
        code, tier = resolve_aar_code("40' Steel Gondola", {})
        assert code == "GS"
        assert tier == "keyword"

    def test_keyword_plain_gondola(self):
        code, tier = resolve_aar_code("Ore Gondola", {})
        assert code == "GB"
        assert tier == "keyword"

    def test_keyword_flat_car(self):
        code, tier = resolve_aar_code("50' Bulkhead Flat Car", {})
        assert code == "FM"
        assert tier == "keyword"

    def test_keyword_tank_car(self):
        code, tier = resolve_aar_code("10,000 Gal Tank Car", {})
        assert code == "TM"
        assert tier == "keyword"

    def test_keyword_reefer(self):
        code, tier = resolve_aar_code("40' Steam Era Ice Bunker Refer", {})
        assert code == "RB"
        assert tier == "keyword"

    def test_keyword_stock_car(self):
        code, tier = resolve_aar_code("40' Wood Stock Car", {})
        assert code == "RS"
        assert tier == "keyword"

    def test_fallback_when_unmatched(self):
        code, tier = resolve_aar_code("Mystery Custom Kitbash", {})
        assert code == "XM"
        assert tier == "fallback"


class TestExtractTonnage:
    def test_extracts_leading_tonnage(self):
        assert extract_tonnage("70 Ton 12 Panel Triple Hopper") == 70

    def test_extracts_tonnage_mid_string(self):
        assert extract_tonnage("52'6 Bethlehem 70-Ton Riveted Drop-End Gondola") == 70

    def test_no_tonnage_returns_none(self):
        assert extract_tonnage("40' Steel SD Box Car") is None


class TestExtractLengthFt:
    def test_extracts_leading_length(self):
        assert extract_length_ft("40' Steel SD Box Car") == 40

    def test_extracts_leading_length_short_form(self):
        assert extract_length_ft("24' Rib Ore Car") == 24

    def test_no_leading_length_returns_none(self):
        assert extract_length_ft("GS 40' Gondola") is None

    def test_strips_leading_whitespace(self):
        assert extract_length_ft("  50' DD Box Car") == 50


class TestResolveCapacityTons:
    def test_regex_extracted_tonnage_used_when_present(self):
        tons, guessed = resolve_capacity_tons("70 Ton 12 Panel Triple Hopper", "HM")
        assert tons == 70
        assert guessed is False

    def test_falls_back_to_default_table(self):
        tons, guessed = resolve_capacity_tons("40' Steel SD Box Car", "XM")
        assert tons == CAPACITY_DEFAULTS["XM"]
        assert guessed is True

    def test_unknown_aar_code_defaults_to_50(self):
        tons, guessed = resolve_capacity_tons("Something Odd", "ZZ")
        assert tons == 50
        assert guessed is True


class TestBuildCars:
    def test_normal_row_imported(self):
        rows = [_row(road="PRR", car_number="12345", type_text="40' Steel SD Box Car")]
        cars, report = build_cars(rows, {})
        assert len(cars) == 1
        assert cars[0].id == "PRR-12345"
        assert cars[0].road == "PRR"
        assert cars[0].car_number == "12345"
        assert cars[0].aar_code == "XM"
        assert cars[0].active is True
        assert report.imported == 1
        assert report.rows_read == 1

    def test_sold_row_skipped(self):
        rows = [_row(inventory_status="SOLD")]
        cars, report = build_cars(rows, {})
        assert cars == []
        assert report.skipped_sold == 1
        assert report.imported == 0

    def test_incomplete_row_skipped(self):
        rows = [_row(road=""), _row(car_number="")]
        cars, report = build_cars(rows, {})
        assert cars == []
        assert report.skipped_incomplete == 2

    def test_duplicate_id_last_row_wins(self):
        rows = [
            _row(road="PRR", car_number="1", type_text="40' Box Car", notes="first"),
            _row(road="PRR", car_number="1", type_text="40' Box Car", notes="second"),
        ]
        cars, report = build_cars(rows, {})
        assert len(cars) == 1
        assert "second" in cars[0].notes
        assert report.duplicate_ids == ["PRR-1"]
        assert report.imported == 1

    def test_map_matched_counted(self):
        rows = [_row(type_text="Custom Decorated Hopper")]
        cmap = {"custom decorated hopper": "LO"}
        _, report = build_cars(rows, cmap)
        assert report.map_matched == 1
        assert report.keyword_matched == 0
        assert report.fallback_matched == 0

    def test_keyword_matched_counted(self):
        rows = [_row(type_text="40' Box Car")]
        _, report = build_cars(rows, {})
        assert report.keyword_matched == 1

    def test_fallback_counted_and_notes_flagged(self):
        rows = [_row(type_text="Mystery Custom Kitbash")]
        cars, report = build_cars(rows, {})
        assert report.fallback_matched == 1
        assert report.unmapped_types == {"Mystery Custom Kitbash": 1}
        assert "guessed aar_code" in cars[0].notes
        assert "Mystery Custom Kitbash" in cars[0].notes

    def test_capacity_regex_vs_default_counted(self):
        rows = [
            _row(type_text="70 Ton 12 Panel Triple Hopper"),
            _row(road="PRR", car_number="2", type_text="40' Box Car"),
        ]
        cars, report = build_cars(rows, {})
        assert report.capacity_regex == 1
        assert report.capacity_default == 1
        guessed_car = next(c for c in cars if c.car_number == "2")
        assert "guessed capacity_tons" in guessed_car.notes

    def test_length_ft_derived_when_present(self):
        rows = [_row(type_text="40' Steel SD Box Car")]
        cars, _ = build_cars(rows, {})
        assert cars[0].length_ft == 40

    def test_length_ft_none_when_absent(self):
        rows = [_row(type_text="GS 40' Gondola")]
        cars, _ = build_cars(rows, {})
        assert cars[0].length_ft is None

    def test_original_notes_preserved_alongside_guess_markers(self):
        rows = [_row(type_text="Mystery Custom Kitbash", notes="Custom lighted")]
        cars, _ = build_cars(rows, {})
        assert "Custom lighted" in cars[0].notes
        assert "guessed aar_code" in cars[0].notes


import yaml as _yaml
from click.testing import CliRunner
from waybill_generator.cli import main


def _make_roster_xlsx(tmp_path, rows: list[dict[str, object]]) -> Path:
    return _make_workbook(tmp_path, rows)


class TestImportRosterCommand:
    def test_writes_cars_yaml(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        source = _make_roster_xlsx(tmp_path, [
            {"A": "45711", "B": "PRR", "F": "40' Steel SD Box Car", "K": "12345", "S": "note"},
        ])

        runner = CliRunner()
        result = runner.invoke(main, [
            "--data-path", str(data_dir),
            "import-roster", "--source", str(source),
        ])

        assert result.exit_code == 0, result.output
        cars = _yaml.safe_load((data_dir / "cars.yaml").read_text())
        assert len(cars) == 1
        assert cars[0]["id"] == "PRR-12345"

    def test_sold_row_excluded_from_output(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        source = _make_roster_xlsx(tmp_path, [
            {"A": "SOLD", "B": "PRR", "F": "40' Box Car", "K": "12345"},
            {"A": "45711", "B": "PRR", "F": "40' Box Car", "K": "99999"},
        ])

        runner = CliRunner()
        result = runner.invoke(main, [
            "--data-path", str(data_dir),
            "import-roster", "--source", str(source),
        ])

        assert result.exit_code == 0, result.output
        cars = _yaml.safe_load((data_dir / "cars.yaml").read_text())
        assert len(cars) == 1
        assert cars[0]["id"] == "PRR-99999"

    def test_report_printed(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        source = _make_roster_xlsx(tmp_path, [
            {"A": "45711", "B": "PRR", "F": "40' Box Car", "K": "12345"},
        ])

        runner = CliRunner()
        result = runner.invoke(main, [
            "--data-path", str(data_dir),
            "import-roster", "--source", str(source),
        ])

        assert "Rows read:" in result.output
        assert "Imported:" in result.output
        assert "AAR code resolution:" in result.output
        assert "Capacity resolution:" in result.output

    def test_full_resync_overwrites_prior_cars_yaml(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        (data_dir / "cars.yaml").write_text(
            "- id: OLD-1\n  road: OLD\n  car_number: '1'\n  aar_code: XM\n  capacity_tons: 50\n"
        )
        source = _make_roster_xlsx(tmp_path, [
            {"B": "PRR", "F": "40' Box Car", "K": "12345"},
        ])

        runner = CliRunner()
        runner.invoke(main, [
            "--data-path", str(data_dir),
            "import-roster", "--source", str(source),
        ])

        cars = _yaml.safe_load((data_dir / "cars.yaml").read_text())
        ids = {c["id"] for c in cars}
        assert "OLD-1" not in ids
        assert "PRR-12345" in ids

    def test_unmapped_types_shown_in_report(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        source = _make_roster_xlsx(tmp_path, [
            {"B": "PRR", "F": "Mystery Custom Kitbash", "K": "12345"},
        ])

        runner = CliRunner()
        result = runner.invoke(main, [
            "--data-path", str(data_dir),
            "import-roster", "--source", str(source),
        ])

        assert "Unmapped types" in result.output
        assert "Mystery Custom Kitbash" in result.output

    def test_uses_car_type_map_from_data_path(self, tmp_path):
        data_dir = tmp_path / "data"
        data_dir.mkdir()
        (data_dir / "car_type_map.yaml").write_text(
            "- source_text: \"Custom Decorated Hopper\"\n  aar_code: LO\n"
        )
        source = _make_roster_xlsx(tmp_path, [
            {"B": "PRR", "F": "Custom Decorated Hopper", "K": "12345"},
        ])

        runner = CliRunner()
        result = runner.invoke(main, [
            "--data-path", str(data_dir),
            "import-roster", "--source", str(source),
        ])

        assert result.exit_code == 0, result.output
        cars = _yaml.safe_load((data_dir / "cars.yaml").read_text())
        assert cars[0]["aar_code"] == "LO"
